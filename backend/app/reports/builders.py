"""Report content (collectors) and file writers (PDF / XLSX / JSON).

A collector returns a neutral dict: title, subtitle, kpis [(label, value)], tables [{name, columns, rows}],
chart {title, x, series[{name, values}]} or None, notes [str]. Writers turn that into a file.
"""
import io
import json
from datetime import datetime

import pandas as pd
from sqlalchemy import case, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.db.models import Anomaly, Recommendation
from app.forecasting.data import models_dir
from app.pipeline import runner
from app.risk import scan

TEMPLATES = {
    "forecast_summary": {"title": "Weekly Demand Forecast Summary", "group": "Forecasting & Revenue",
                         "description": "Consolidated 30-day forward demand by category and store, top SKUs and model accuracy."},
    "risk_dossier": {"title": "Risk & Forensic Audit Dossier", "group": "Audit & Compliance",
                     "description": "Open anomalies by type and severity with the highest-exposure alerts and detector quality."},
    "replenishment_plan": {"title": "Store Replenishment & Inventory Plan", "group": "Supply Chain & Ops",
                           "description": "Prescriptive reorders, inter-store transfers and markdowns awaiting or past approval."},
}


def _inr(v) -> str:
    return f"Rs {v:,.0f}"


def _names():
    base = settings.data_dir / "processed"
    return (pd.read_csv(base / "stores.csv").set_index("store")["name"].to_dict(),
            pd.read_csv(base / "items.csv").set_index("item")[["name", "category"]].to_dict("index"))


def collect_forecast_summary(db: Session) -> dict:
    fc = pd.read_parquet(models_dir() / "forecasts.parquet")
    hist = pd.read_parquet(runner.paths()["features"], columns=["date", "store", "item", "sales"])
    stores, items = _names()
    last = hist["date"].max()
    cur = hist[hist["date"] > last - pd.Timedelta(days=30)]
    f30 = fc[fc["date"] < fc["date"].min() + pd.Timedelta(days=30)]
    meta = json.loads((models_dir() / "benchmark.json").read_text(encoding="utf-8"))
    serving = next(m for m in meta["models"] if m["name"] == meta["serving"])
    cat = lambda i: items[i]["category"]  # noqa: E731

    def table(group_cur, group_fc, name_map=None):
        a, f = cur.groupby(group_cur)["sales"].sum(), f30.groupby(group_fc)["yhat"].sum()
        df = pd.DataFrame({"last30": a, "next30": f}).fillna(0)
        df["growth"] = (df["next30"] / df["last30"] - 1) * 100
        return [[name_map.get(k, k) if name_map else k, f"{r.last30:,.0f}", f"{r.next30:,.0f}", f"{r.growth:+.1f}%"] for k, r in df.sort_values("next30", ascending=False).iterrows()]

    top = f30.groupby("item")["yhat"].sum().sort_values(ascending=False).head(10)
    chain_h = hist[hist["date"] > last - pd.Timedelta(days=60)].groupby("date")["sales"].sum()
    chain_f = fc[fc["date"] < fc["date"].min() + pd.Timedelta(days=30)].groupby("date")["yhat"].sum()
    return {
        "title": TEMPLATES["forecast_summary"]["title"], "subtitle": f"Forecast window {meta['forecast_window'][0]} to {meta['forecast_window'][1]} · model {meta['serving']} {meta['version']}",
        "kpis": [("Forecast next 30 days (units)", f"{f30['yhat'].sum():,.0f}"), ("Actual last 30 days (units)", f"{cur['sales'].sum():,.0f}"),
                 ("Series forecast", f"{fc[['store', 'item']].drop_duplicates().shape[0]}"), ("Model MAPE / R2", f"{serving['full_set']['mape']:.1f}% / {serving['full_set']['r2']:.2f}")],
        "tables": [
            {"name": "By category", "columns": ["Category", "Last 30d units", "Next 30d forecast", "Change"], "rows": table(cur["item"].map(cat), f30["item"].map(cat))},
            {"name": "By store", "columns": ["Store", "Last 30d units", "Next 30d forecast", "Change"], "rows": table(cur["store"], f30["store"], stores)},
            {"name": "Top 10 SKUs by forecast demand", "columns": ["SKU", "Category", "Next 30d forecast"], "rows": [[items[i]["name"], items[i]["category"], f"{v:,.0f}"] for i, v in top.items()]},
        ],
        "chart": {"title": "Chain-wide daily units: last 60 days actual and next 30 days forecast",
                  "x": [d.strftime("%d %b") for d in list(chain_h.index) + list(chain_f.index)],
                  "series": [{"name": "Actual", "values": [float(v) for v in chain_h.values] + [None] * len(chain_f)},
                             {"name": "Forecast", "values": [None] * len(chain_h) + [float(v) for v in chain_f.values]}]},
        "notes": ["Forecasts are decision aids with the uncertainty shown on the Forecasting page; the model was backtested on a 90-day holdout."],
    }


def collect_risk_dossier(db: Session) -> dict:
    stores, items = _names()
    sev = case({"High": 0, "Medium": 1, "Low": 2}, value=Anomaly.severity, else_=3)
    rows = db.scalars(select(Anomaly).where(Anomaly.status.in_(["Open", "In Review"])).order_by(sev, Anomaly.amount.desc())).all()
    df = pd.DataFrame([{"type": r.anomaly_type, "severity": r.severity, "amount": r.amount} for r in rows])
    m = scan.read_metrics() or {}
    ens = (m.get("detectors") or {}).get("Ensemble", {})
    by_type = [[t.replace("_", " "), str(len(g)), _inr(g["amount"].sum()), str((g["severity"] == "High").sum())] for t, g in df.groupby("type")] if len(df) else []
    return {
        "title": TEMPLATES["risk_dossier"]["title"], "subtitle": f"Anomaly scan {m.get('ran_at', 'not run')} · {m.get('rows_scanned', 0):,} billed rows scanned",
        "kpis": [("Open alerts", f"{len(rows)}"), ("High severity", f"{int((df['severity'] == 'High').sum()) if len(df) else 0}"),
                 ("High-severity exposure", _inr(df.loc[df['severity'] == 'High', 'amount'].sum()) if len(df) else "Rs 0"),
                 ("Ensemble precision / recall", f"{ens.get('precision', 0):.2f} / {ens.get('recall', 0):.2f}")],
        "tables": [
            {"name": "Open alerts by type", "columns": ["Type", "Alerts", "Exposure", "High severity"], "rows": by_type},
            {"name": "Top 25 alerts by severity and exposure", "columns": ["Alert", "Type", "Store", "Date", "Exposure", "Severity", "Status", "Why flagged"],
             "rows": [[f"ALT-{r.id:04d}", r.anomaly_type.replace("_", " "), stores.get(r.store, ""), r.date.isoformat(), _inr(r.amount), r.severity, r.status, r.explanation or ""] for r in rows[:25]]},
            {"name": "Detector quality (vs injected labels)", "columns": ["Detector", "Precision", "Recall", "F1"],
             "rows": [[k, f"{v['precision']:.2f}", f"{v['recall']:.2f}", f"{v['f1']:.2f}"] for k, v in (m.get("detectors") or {}).items()]},
        ],
        "chart": None,
        "notes": ["Exposure is an estimate: excess billed value, below-cost margin loss, or lost sales over the affected days.",
                  "Detector quality is measured on synthetic injected anomalies and understates real precision."],
    }


def collect_replenishment_plan(db: Session) -> dict:
    stores, items = _names()
    recs = db.scalars(select(Recommendation).where(Recommendation.status.in_(["Pending", "Approved"])).order_by(Recommendation.impact_inr.desc())).all()

    def rows(kind):
        out = []
        for r in (x for x in recs if x.kind == kind):
            pl = r.payload or {}
            out.append([r.id, r.status, stores.get(r.store, ""), stores.get(r.store_to, "") if r.store_to else "", pl.get("sku", ""),
                        pl.get("qty", pl.get("discount_pct", "")), _inr(r.impact_inr), f"{r.confidence:.0%}", r.title])
        return out

    cols = ["ID", "Status", "Store", "To store", "SKU", "Qty / %", "Impact", "Confidence", "Action"]
    pend = [r for r in recs if r.status == "Pending"]
    return {
        "title": TEMPLATES["replenishment_plan"]["title"], "subtitle": f"{len(recs)} pending or approved actions · generated from the latest forecast and stock position",
        "kpis": [("Reorders", str(sum(r.kind == "reorder" for r in recs))), ("Transfers", str(sum(r.kind == "transfer" for r in recs))),
                 ("Markdowns", str(sum(r.kind == "markdown" for r in recs))), ("Pending value", _inr(sum(r.impact_inr for r in pend)))],
        "tables": [{"name": "Purchase orders (reorders)", "columns": cols, "rows": rows("reorder")},
                   {"name": "Inter-store transfers", "columns": cols, "rows": rows("transfer")},
                   {"name": "Markdowns", "columns": cols, "rows": rows("markdown")}],
        "chart": None,
        "notes": ["Impact figures mix measures (protected revenue, capital released, margin gain) as labelled on the Recommendations page; they are estimates.",
                  "Lead time, service level, elasticities and write-off rates are assumptions - see the calculation audit trail of each action."],
    }


COLLECTORS = {"forecast_summary": collect_forecast_summary, "risk_dossier": collect_risk_dossier, "replenishment_plan": collect_replenishment_plan}


# ------------------------------------------------------------------ writers
def write_json(content: dict) -> bytes:
    return json.dumps({**content, "generated_at": datetime.utcnow().isoformat(timespec="seconds")}, indent=1, default=str).encode("utf-8")


def write_xlsx(content: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append([content["title"]])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([content["subtitle"]])
    ws.append([])
    for k, v in content["kpis"]:
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    for n in content["notes"]:
        ws.append([])
        ws.append([n])
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 28
    head = PatternFill("solid", fgColor="1E294B")
    for t in content["tables"]:
        sheet = wb.create_sheet(t["name"][:31].replace("/", "-"))
        sheet.append(t["columns"])
        for c in sheet[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), head
        for r in t["rows"]:
            sheet.append(r)
        for i, col in enumerate(t["columns"], 1):
            width = max([len(str(col))] + [len(str(r[i - 1])) for r in t["rows"][:200]]) + 2
            sheet.column_dimensions[get_column_letter(i)].width = min(width, 70)
        sheet.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def write_pdf(content: dict) -> bytes:
    from reportlab.graphics.charts.lineplots import LinePlot
    from reportlab.graphics.shapes import Drawing
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    ss = getSampleStyleSheet()
    cell = ParagraphStyle("cell", parent=ss["BodyText"], fontSize=7.5, leading=9)
    head = ParagraphStyle("head", parent=cell, textColor=colors.white, fontName="Helvetica-Bold")
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm, bottomMargin=14 * mm, title=content["title"])
    el = [Paragraph(content["title"], ss["Title"]), Paragraph(content["subtitle"], ss["Normal"]), Spacer(1, 6 * mm)]
    k = Table([[Paragraph(f"<b>{v}</b><br/><font size=7 color='#555555'>{l}</font>", ss["BodyText"]) for l, v in content["kpis"]]], hAlign="LEFT")
    k.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.5, colors.grey), ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.lightgrey), ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    el += [k, Spacer(1, 5 * mm)]
    ch = content.get("chart")
    if ch and len(ch["x"]) > 1:
        el.append(Paragraph(ch["title"], ss["Heading4"]))
        d = Drawing(250 * mm, 62 * mm)
        lp = LinePlot()
        lp.x, lp.y, lp.width, lp.height = 12 * mm, 8 * mm, 225 * mm, 48 * mm
        lp.data = [[(i, v) for i, v in enumerate(s["values"]) if v is not None] for s in ch["series"]]
        for i, col in enumerate([colors.HexColor("#00A1DA"), colors.HexColor("#8B7CF6")]):
            if i < len(lp.lines):
                lp.lines[i].strokeColor, lp.lines[i].strokeWidth = col, 1.4
        lp.xValueAxis.visible = False
        lp.yValueAxis.labels.fontSize = 7
        d.add(lp)
        el += [d, Paragraph("<font color='#00A1DA'>■</font> Actual &nbsp; <font color='#8B7CF6'>■</font> Forecast", ss["Normal"]), Spacer(1, 4 * mm)]
    avail = 269 * mm
    for t in content["tables"]:
        el.append(Paragraph(t["name"], ss["Heading3"]))
        if not t["rows"]:
            el += [Paragraph("No rows.", ss["Italic"]), Spacer(1, 3 * mm)]
            continue
        wide = [i for i, c in enumerate(t["columns"]) if c in ("Why flagged", "Action")]
        widths = [(avail * 0.4 if i in wide else None) for i in range(len(t["columns"]))]
        rest = (avail - sum(w for w in widths if w)) / max(1, sum(w is None for w in widths))
        widths = [w or rest for w in widths]
        data = [[Paragraph(str(c), head) for c in t["columns"]]] + [[Paragraph(str(c), cell) for c in r] for r in t["rows"][:40]]
        tb = Table(data, colWidths=widths, repeatRows=1)
        tb.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E294B")), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F6FF")]),
                                ("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        el.append(tb)
        if len(t["rows"]) > 40:
            el.append(Paragraph(f"Showing 40 of {len(t['rows'])} rows - the XLSX export contains all rows.", ss["Italic"]))
        el.append(Spacer(1, 4 * mm))
    for n in content["notes"]:
        el.append(Paragraph(f"• {n}", ss["Italic"]))

    def footer(canvas, d):
        canvas.setFont("Helvetica", 7)
        canvas.drawString(14 * mm, 8 * mm, f"DecisionIQ · generated {datetime.utcnow():%Y-%m-%d %H:%M} UTC · decision-support estimates, final decisions rest with business managers")
        canvas.drawRightString(283 * mm, 8 * mm, f"Page {d.page}")

    doc.build(el, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


WRITERS = {"pdf": (write_pdf, "application/pdf"), "xlsx": (write_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
           "json": (write_json, "application/json")}
