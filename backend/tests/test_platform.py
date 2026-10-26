import io
import json
from datetime import date

import numpy as np
import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.core.config import settings
from app.db.models import Anomaly
from app.db.session import get_engine
from app.main import app
from tests.test_recommend import LAST, make_state

client = TestClient(app)


@pytest.fixture
def world(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "data_dir", tmp_path)
    proc = tmp_path / "processed"
    (proc / "models").mkdir(parents=True)
    st = make_state()
    st["fc"].to_parquet(proc / "models" / "forecasts.parquet")
    st["items"].to_csv(proc / "items.csv", index=False)
    st["stores"].to_csv(proc / "stores.csv", index=False)
    hist = st["hist"].merge(st["cur"][["store", "item", "on_hand", "on_order", "list_price"]], on=["store", "item"])
    hist.to_csv(proc / "sales.csv", index=False)
    hist[["date", "store", "item", "sales"]].to_parquet(proc / "features.parquet")
    meta = {"serving": "XGBoost", "version": "xgboost-test", "trained_at": "2026-01-01T00:00:00", "forecast_window": ["2018-01-01", "2018-03-31"],
            "models": [{"name": "XGBoost", "full_set": {"mape": 12.0, "r2": 0.9, "wape": 10.0}}]}
    (proc / "models" / "benchmark.json").write_text(json.dumps(meta))
    return tmp_path


# ---------------------------------------------------------------- settings & audit
def test_settings_defaults_update_validation_and_reset():
    s = client.get("/api/settings").json()
    assert s["config"]["risk"]["z_threshold"] == 3.0 and s["user"]["role"] == "Admin"
    assert "z_threshold" in s["applied"]["risk"] and len(s["rbac"]) == 7

    r = client.put("/api/settings/risk", json={"z_threshold": 3.5}).json()
    assert r["changed"] == ["z_threshold"] and client.get("/api/settings").json()["config"]["risk"]["z_threshold"] == 3.5
    assert client.put("/api/settings/risk", json={"z_threshold": 99}).status_code == 422
    assert client.put("/api/settings/nope", json={}).status_code == 404

    client.post("/api/settings/reset")
    assert client.get("/api/settings").json()["config"]["risk"]["z_threshold"] == 3.0


def test_audit_trail_is_hash_chained_and_detects_tampering():
    client.put("/api/settings/inventory", json={"lead_time_days": 5})
    client.put("/api/settings/inventory", json={"lead_time_days": 6})
    log = client.get("/api/audit").json()
    assert log["total"] == 2 and "lead_time_days: 5 -> 6" in log["entries"][0]["action"]
    assert client.get("/api/audit/verify").json()["ok"] is True

    with Session(get_engine()) as db:  # rewrite history behind the application's back
        db.execute(text("UPDATE audit_log SET action = 'nothing happened' WHERE id = 1"))
        db.commit()
    v = client.get("/api/audit/verify").json()
    assert v["ok"] is False and v["broken_at"] == 1


def test_inventory_settings_drive_recommendations(world):
    client.post("/api/recommendations/generate")
    md = lambda: [(r["store"], r["item"]) for r in client.get("/api/recommendations?kind=markdown").json()["recommendations"]]  # noqa: E731
    assert (1, 2) in md()  # 100 days of cover is overstock at the default 45-day threshold
    client.put("/api/settings/inventory", json={"overstock_cover_days": 120})
    client.post("/api/recommendations/generate")
    assert (1, 2) not in md()
    assert any(e["category"] == "recommendation" for e in client.get("/api/audit").json()["entries"])


# ---------------------------------------------------------------- reports
def seed_alerts():
    engine = get_engine()
    Anomaly.__table__.create(engine, checkfirst=True)
    with Session(engine) as db:
        db.add_all([Anomaly(store=1, item=1, date=date(2017, 12, 1), end_date=date(2017, 12, 1), anomaly_type="duplicate_billing", score=0.99,
                            severity="High", amount=86400, explanation="Billed 42 vs 21 expected"),
                    Anomaly(store=2, item=None, date=date(2017, 12, 4), anomaly_type="cost_spike", score=0.7, severity="Medium", amount=14000)])
        db.commit()


@pytest.mark.parametrize("template", ["forecast_summary", "risk_dossier", "replenishment_plan"])
def test_reports_all_templates_and_formats(world, template):
    seed_alerts()
    client.post("/api/recommendations/generate")
    for fmt, magic in (("pdf", b"%PDF"), ("xlsx", b"PK"), ("json", b"{")):
        r = client.post("/api/reports/generate", json={"template": template, "format": fmt})
        assert r.status_code == 201, r.text
        rep = r.json()
        body = client.get(f"/api/reports/{rep['id']}/download").content
        assert body.startswith(magic) and len(body) == rep["size_bytes"] > 500
        import hashlib

        assert hashlib.sha256(body).hexdigest() == rep["sha256"]
        if fmt == "xlsx":
            wb = load_workbook(io.BytesIO(body))
            assert wb.sheetnames[0] == "Summary" and len(wb.sheetnames) >= 3
        if fmt == "json":
            assert json.loads(body)["title"]


def test_reports_ledger_templates_and_errors(world):
    seed_alerts()
    client.post("/api/reports/generate", json={"template": "risk_dossier", "format": "json"})
    client.post("/api/reports/generate", json={"template": "forecast_summary", "format": "json"})
    led = client.get("/api/reports").json()
    assert led["total"] == 2 and led["storage_bytes"] > 0 and led["reports"][0]["template"] == "forecast_summary"
    assert client.get("/api/reports?template=risk_dossier").json()["total"] == 1
    assert client.get("/api/reports?q=Forensic").json()["total"] == 1
    t = {x["key"]: x for x in client.get("/api/reports/templates").json()["templates"]}
    assert t["risk_dossier"]["runs"] == 1 and t["replenishment_plan"]["last_run"] is None
    assert client.post("/api/reports/generate", json={"template": "bogus"}).status_code == 422
    assert client.get("/api/reports/999/download").status_code == 404
    assert any("Generated report" in e["action"] for e in client.get("/api/audit").json()["entries"])


def test_report_without_data_is_a_clear_409(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "data_dir", tmp_path)
    assert client.post("/api/reports/generate", json={"template": "forecast_summary", "format": "pdf"}).status_code == 409


# ---------------------------------------------------------------- dashboard & search
def test_dashboard_summary_status_and_search(world):
    seed_alerts()
    client.post("/api/recommendations/generate")
    d = client.get("/api/dashboard/summary").json()
    assert d["ready"] and d["kpis"]["units_30d"] > 0 and d["regions"] and d["attention"][0]["severity"] == "High"
    assert d["recommendations"] and d["inventory"]["total_skus"] == 6 and d["open_alerts"] == 2
    st = client.get("/api/dashboard/status").json()
    assert [s["key"] for s in st["steps"]] == ["data", "clean", "models", "forecast", "risk"] and st["pending_recs"] > 0 and st["high_alerts"] == 1
    hits = client.get("/api/search?q=blr").json()["results"]
    assert any(h["type"] == "Store" for h in hits)
    assert client.get("/api/search?q=alt-1").json()["results"][0]["type"] == "Alert"
    assert client.get("/api/search?q=x").json()["results"] == []


def test_dashboard_not_ready_without_data(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "data_dir", tmp_path)
    assert client.get("/api/dashboard/summary").json() == {"ready": False}
    assert client.get("/api/dashboard/status").json()["healthy"] is False
